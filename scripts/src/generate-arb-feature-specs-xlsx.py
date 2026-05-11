#!/usr/bin/env python3
"""Generate an Excel workbook of CareOps360 functional specs (one tab per module + flat + index)."""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "out"
OUT_DIR.mkdir(parents=True, exist_ok=True)
XLSX_PATH = OUT_DIR / "CareOps360_Feature_Specifications.xlsx"

# Load FEATURE_SPECS from the docx generator script (single source of truth).
DOC_GEN = ROOT / "scripts" / "src" / "generate-arb-doc.py"
spec = importlib.util.spec_from_file_location("arb_doc", DOC_GEN)
mod = importlib.util.module_from_spec(spec)
sys.modules["arb_doc"] = mod
spec.loader.exec_module(mod)
FEATURE_SPECS = mod.FEATURE_SPECS

# Styling -------------------------------------------------------------
NAVY = "FF232F3E"
ORANGE = "FFFF9900"
LIGHT = "FFF2F3F3"
SUBHEAD = "FFE6ECF2"

THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FONT = Font(name="Calibri", size=11, bold=True, color=NAVY)
SUB_FILL = PatternFill("solid", fgColor=SUBHEAD)
BODY_FONT = Font(name="Calibri", size=11)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_col_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_header_row(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def style_body(ws, start_row: int, end_row: int, cols: int) -> None:
    for r in range(start_row, end_row + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER


# Sheet builders ------------------------------------------------------
ATTRIBUTES = [
    ("Purpose", lambda fs: [fs["purpose"]]),
    ("Primary Users / Roles", lambda fs: fs["users"]),
    ("Key Capabilities", lambda fs: fs["capabilities"]),
    ("KPIs & Visualizations", lambda fs: fs["kpis"]),
    ("Data Inputs (Near-Term)", lambda fs: fs["data_inputs"]),
    ("Interactions & Admin Actions", lambda fs: fs["interactions"]),
    ("Acceptance Criteria", lambda fs: fs["acceptance"]),
    ("Future Enhancements (Long-Term)", lambda fs: fs["future"]),
]


def build_index(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Index"
    ws["A1"] = "CareOps360 — Functional Specifications"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    ws["A2"] = "Horizon BCBS Payer Command Center  |  ARB companion workbook"
    ws["A2"].font = Font(name="Calibri", size=11, italic=True, color="FF555555")
    ws.merge_cells("A1:E1")
    ws.merge_cells("A2:E2")

    headers = ["#", "Module", "Purpose (one-liner)", "Primary Users", "Sheet"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=h)
    style_header_row(ws, 4, len(headers))

    for i, fs in enumerate(FEATURE_SPECS, start=1):
        row = 4 + i
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=fs["title"])
        ws.cell(row=row, column=3, value=fs["purpose"])
        ws.cell(row=row, column=4, value=", ".join(fs["users"]))
        sheet_name = sheet_name_for(fs["title"])
        link_cell = ws.cell(row=row, column=5, value=sheet_name)
        link_cell.hyperlink = f"#'{sheet_name}'!A1"
        link_cell.font = Font(name="Calibri", size=11, color="FF0563C1", underline="single")

    style_body(ws, 5, 4 + len(FEATURE_SPECS), len(headers))
    set_col_widths(ws, [4, 32, 70, 36, 28])


def sheet_name_for(title: str) -> str:
    # Excel sheet names: max 31 chars, no : \ / ? * [ ]
    safe = title.replace("/", "-").replace(":", "-").replace("?", "")
    return safe[:31]


def build_module_sheet(wb: Workbook, fs: dict) -> None:
    ws = wb.create_sheet(sheet_name_for(fs["title"]))
    ws["A1"] = fs["title"]
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    ws["A2"] = "Functional specification"
    ws["A2"].font = Font(name="Calibri", size=11, italic=True, color="FF555555")
    ws.merge_cells("A1:C1")
    ws.merge_cells("A2:C2")

    # Header
    headers = ["Section", "Item", "Notes / Owner / Status"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=h)
    style_header_row(ws, 4, len(headers))

    # Body
    row = 5
    first_body = row
    for section_name, getter in ATTRIBUTES:
        items = getter(fs)
        section_start = row
        for item in items:
            ws.cell(row=row, column=1, value=section_name)
            ws.cell(row=row, column=2, value=item)
            ws.cell(row=row, column=3, value="")  # blank for team to populate
            row += 1
        # Merge section column for the group
        if row - 1 > section_start:
            ws.merge_cells(start_row=section_start, start_column=1,
                           end_row=row - 1, end_column=1)
        # Apply subhead fill on the section column for the merged block
        cell = ws.cell(row=section_start, column=1)
        cell.fill = SUB_FILL
        cell.font = SUB_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

    style_body(ws, first_body, row - 1, len(headers))
    # Re-apply subhead styling after style_body overwrote fonts on merged cells
    r = first_body
    for section_name, getter in ATTRIBUTES:
        items = getter(fs)
        if not items:
            continue
        cell = ws.cell(row=r, column=1)
        cell.fill = SUB_FILL
        cell.font = SUB_FONT
        r += len(items)

    set_col_widths(ws, [30, 80, 36])
    ws.freeze_panes = "A5"


def build_flat_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("All Specs (Flat)")
    headers = ["Module", "Section", "Item", "Owner", "Status", "Target Date", "Notes"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, 1, len(headers))

    row = 2
    for fs in FEATURE_SPECS:
        for section_name, getter in ATTRIBUTES:
            for item in getter(fs):
                ws.cell(row=row, column=1, value=fs["title"])
                ws.cell(row=row, column=2, value=section_name)
                ws.cell(row=row, column=3, value=item)
                # owner / status / date / notes intentionally blank
                row += 1

    style_body(ws, 2, row - 1, len(headers))
    set_col_widths(ws, [28, 30, 80, 22, 16, 16, 30])
    ws.auto_filter.ref = f"A1:G{row - 1}"


def build_acceptance_sheet(wb: Workbook) -> None:
    """A focused checklist of just acceptance criteria across all modules."""
    ws = wb.create_sheet("Acceptance Checklist")
    headers = ["Module", "Acceptance Criterion", "Pass / Fail", "Tested By", "Tested On", "Notes"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, 1, len(headers))

    row = 2
    for fs in FEATURE_SPECS:
        for criterion in fs["acceptance"]:
            ws.cell(row=row, column=1, value=fs["title"])
            ws.cell(row=row, column=2, value=criterion)
            row += 1
    style_body(ws, 2, row - 1, len(headers))
    set_col_widths(ws, [28, 80, 14, 22, 16, 30])
    ws.auto_filter.ref = f"A1:F{row - 1}"


# Build ---------------------------------------------------------------
def main() -> None:
    wb = Workbook()
    build_index(wb)
    for fs in FEATURE_SPECS:
        build_module_sheet(wb, fs)
    build_flat_sheet(wb)
    build_acceptance_sheet(wb)
    wb.save(XLSX_PATH)
    print(f"[arb-xlsx] wrote {XLSX_PATH}")
    print(f"[arb-xlsx] sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
